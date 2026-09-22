# -*- coding: utf-8 -*-
"""保護セル（入力できない欄）チェッカー

申請書を作る前に必ず通す。書き込もうとしているセルが
「Excelで開いても入力できない＝保護されている欄」でないかを判定する。

保護セルは中身が数式（合計＝SUMなど）であることが多く、
プログラムから書き込むと保護を無視して書けてしまい、数式を壊す。

使い方：
  1) ファイル全体の保護状況を見る
     python3 check_locked.py <ファイル.xlsx>

  2) 書き込み予定のセルが入力可能か事前判定する
     python3 check_locked.py <ファイル.xlsx> "申請書1枚目!CX32" "申請書1枚目!CX29"

判定結果：
  ✅ 入力可   … 書き込んでよい
  🚫 入力不可 … 書かない。空欄のまま残し、社長に報告する
"""
import sys
import openpyxl
import warnings
warnings.filterwarnings("ignore")


def check(path, targets=None):
    wb = openpyxl.load_workbook(path)

    if targets:
        print(f"■ 書き込み可否チェック：{path}\n")
        ng = []
        for t in targets:
            if "!" in t:
                sn, coord = t.split("!", 1)
            else:
                sn, coord = wb.sheetnames[0], t
            ws = wb[sn]
            cell = ws[coord]
            protected = ws.protection.sheet and cell.protection.locked
            mark = "🚫 入力不可" if protected else "✅ 入力可"
            formula = ""
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula = f"  ※自動計算の数式: {cell.value}"
            print(f"  {mark}  {sn}!{coord}{formula}")
            if protected:
                ng.append(f"{sn}!{coord}")
        if ng:
            print(f"\n⚠️ 次の欄は入力できません。書き込まず空欄のままにし、社長に報告してください：")
            for x in ng:
                print(f"   ・{x}")
        return 1 if ng else 0

    # ファイル全体の概観
    print(f"■ シート保護と自動計算欄の一覧：{path}\n")
    for ws in wb.worksheets:
        state = "保護あり" if ws.protection.sheet else "保護なし（全セル入力可）"
        print(f"[{ws.title}] {state}")
        if not ws.protection.sheet:
            continue
        formulas = []
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas.append((c.coordinate, c.value))
        if formulas:
            print(f"   自動計算の数式 {len(formulas)}件（ここには書き込まない）:")
            for coord, f in formulas[:15]:
                print(f"     🔒 {coord}: {f[:70]}")
            if len(formulas) > 15:
                print(f"     … 他 {len(formulas)-15} 件")
        print()
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(check(sys.argv[1], sys.argv[2:] or None))
